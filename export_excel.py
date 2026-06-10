"""
Excel Report Generator
Creates a formatted multi-sheet .xlsx workbook for investor reporting,
demonstrating Excel/PowerBI-adjacent data skills.
"""

import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

from config import get_esg_dataframe, get_emissions_trajectory_df
from sfdr_pai import generate_pai_report

OUTPUT = "data/InfraESG_Report.xlsx"

HEADER_FONT = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
TITLE_FONT = Font(name="Calibri", bold=True, size=14, color="2E7D32")
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


def _style_header(ws, cols: int, row: int = 1):
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER


def _auto_width(ws):
    for col_cells in ws.columns:
        max_len = 0
        col_letter = None
        for cell in col_cells:
            if col_letter is None and hasattr(cell, "column_letter"):
                col_letter = cell.column_letter
            try:
                val = str(cell.value) if cell.value is not None else ""
                max_len = max(max_len, len(val))
            except Exception:
                pass
        if col_letter:
            ws.column_dimensions[col_letter].width = min(max_len + 3, 40)


def _write_df(ws, df: pd.DataFrame, start_row: int = 1) -> int:
    """Write DataFrame to worksheet and return the next empty row."""
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), start_row):
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center")
    _style_header(ws, len(df.columns), start_row)
    return start_row + len(df) + 1


def generate_report():
    esg = get_esg_dataframe()
    trajectory = get_emissions_trajectory_df()
    pai_report = generate_pai_report(esg)

    wb = Workbook()

    # ── Sheet 1: Executive Summary ────────────────────────────────────────
    ws = wb.active
    ws.title = "Executive Summary"
    ws.cell(row=1, column=1, value="InfraESG Analytics — Portfolio Report").font = TITLE_FONT
    ws.merge_cells("A1:E1")

    summary_data = [
        ("Total Companies", len(esg)),
        ("Sectors Covered", esg["sector"].nunique()),
        ("Average ESG Score", round(esg["esg_score"].mean(), 1)),
        ("Avg Carbon Intensity (tCO₂e/€M rev)", round(esg["carbon_intensity"].mean(), 1)),
        ("EU Taxonomy Aligned (%)", pai_report["taxonomy"]["portfolio_aligned_pct"]),
        ("SBTi-Validated Companies", pai_report["taxonomy"]["sbti_validated_count"]),
        ("SFDR Red Flags", pai_report["summary"]["red_flags"]),
        ("SFDR Data Gaps (disclosed)", pai_report["summary"]["data_gaps"]),
    ]
    for i, (label, val) in enumerate(summary_data, 3):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        ws.cell(row=i, column=2, value=val)

    _auto_width(ws)

    # ── Sheet 2: ESG Scores ───────────────────────────────────────────────
    ws2 = wb.create_sheet("ESG Scores")
    score_cols = [
        "name", "ticker", "sector", "country",
        "esg_score", "env_score", "social_score", "gov_score",
        "carbon_intensity", "renewable_energy_pct",
    ]
    df2 = esg[score_cols].sort_values("esg_score", ascending=False).copy()
    df2["carbon_intensity"] = df2["carbon_intensity"].round(1)
    _write_df(ws2, df2)
    _auto_width(ws2)

    # ── Sheet 3: SFDR PAI Indicators ──────────────────────────────────────
    ws3 = wb.create_sheet("SFDR PAI")
    pai_df = pai_report["pai_indicators"][
        ["pai_id", "indicator", "metric", "value", "unit", "status"]
    ]
    _write_df(ws3, pai_df)
    _auto_width(ws3)

    # ── Sheet 4: EU Taxonomy ──────────────────────────────────────────────
    ws4 = wb.create_sheet("EU Taxonomy")
    tax_cols = [
        "name", "ticker", "sector",
        "taxonomy_eligible_pct", "taxonomy_aligned_pct",
        "has_sbti", "net_zero_year",
    ]
    _write_df(ws4, esg[tax_cols].sort_values("taxonomy_aligned_pct", ascending=False))
    _auto_width(ws4)

    # ── Sheet 5: Decarbonisation ──────────────────────────────────────────
    ws5 = wb.create_sheet("Decarbonisation")
    pivot = trajectory.pivot(index="ticker", columns="year", values="emissions_ktco2e")
    year_cols = sorted(pivot.columns)
    first_year, last_year = str(year_cols[0]), str(year_cols[-1])
    pivot.columns = [str(c) for c in pivot.columns]
    pivot = pivot.reset_index()
    merged = esg[["ticker", "name", "sector"]].merge(pivot, on="ticker")
    merged["reduction_pct"] = round(
        (1 - merged[last_year] / merged[first_year]) * 100, 1
    )
    _write_df(ws5, merged.sort_values("reduction_pct", ascending=False))
    _auto_width(ws5)

    # ── Sheet 6: Governance & Social ──────────────────────────────────────
    ws6 = wb.create_sheet("Governance & Social")
    gov_cols = [
        "name", "ticker", "employees",
        "women_mgmt_pct", "board_diversity_pct", "board_independence_pct",
        "controversy_score",
    ]
    _write_df(ws6, esg[gov_cols].sort_values("board_diversity_pct", ascending=False))
    _auto_width(ws6)

    # ── Save ──────────────────────────────────────────────────────────────
    os.makedirs(os.path.dirname(OUTPUT), exist_ok=True)
    wb.save(OUTPUT)
    print(f"Report saved to {OUTPUT}")
    return OUTPUT


if __name__ == "__main__":
    generate_report()
